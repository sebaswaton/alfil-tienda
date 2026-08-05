from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
from typing import Iterable

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app import models, schemas
from app.product_creation import (
    audit_product_creation,
    create_product_record,
    normalize_sku,
)
from app.product_excel_template import (
    PRODUCT_CONDITION_VALUES,
    PRODUCT_CURRENCY_VALUES,
    PRODUCT_HEADERS,
    PRODUCT_STATUS_VALUES,
)


MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_PRODUCT_ROW = 300
STOCK_NOTE_DEFAULT = "Disponible bajo cotización"
STOCK_TYPE_DEFAULT = "serialized"

CONDITION_VALUES = {
    PRODUCT_CONDITION_VALUES[0].casefold(): False,
    PRODUCT_CONDITION_VALUES[1].casefold(): True,
}
CURRENCY_VALUES = {value.casefold(): value for value in PRODUCT_CURRENCY_VALUES}
STATUS_VALUES = {
    PRODUCT_STATUS_VALUES[0].casefold(): "active",
    PRODUCT_STATUS_VALUES[1].casefold(): "draft",
    PRODUCT_STATUS_VALUES[2].casefold(): "inactive",
}


class ProductImportFileError(ValueError):
    pass


def safe_upload_filename(filename: str | None) -> str:
    name = re.split(r"[\\/]", filename or "archivo.xlsx")[-1]
    cleaned = "".join(character for character in name if character.isprintable()).strip()
    return cleaned[:255] or "archivo.xlsx"


def _is_empty(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _json_value(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _text_value(
    value,
    *,
    label: str,
    errors: list[str],
    required: bool = False,
    min_length: int | None = None,
    max_length: int | None = None,
) -> str:
    if _is_empty(value):
        if required:
            errors.append(f"{label} es obligatorio.")
        return ""
    if not isinstance(value, str):
        errors.append(f"{label} debe ser texto.")
        return ""
    normalized = value.strip()
    if min_length is not None and len(normalized) < min_length:
        errors.append(f"{label} debe tener al menos {min_length} caracteres.")
    if max_length is not None and len(normalized) > max_length:
        errors.append(f"{label} no puede superar {max_length} caracteres.")
    return normalized


def _build_name_index(items: Iterable[tuple[int, str]]):
    index: dict[str, list[tuple[int, str]]] = defaultdict(list)
    for item_id, name in items:
        index[name.strip().casefold()].append((item_id, name))
    return index


def _resolve_catalog_name(
    value,
    *,
    label: str,
    index: dict[str, list[tuple[int, str]]],
    errors: list[str],
) -> tuple[int | None, str]:
    text = _text_value(value, label=label, errors=errors, required=True)
    if not text:
        return None, ""
    matches = index.get(text.casefold(), [])
    if not matches:
        errors.append(f"{label} no existe en el catálogo.")
        return None, text
    if len(matches) > 1:
        errors.append(f"{label} es ambiguo; existen varios registros con ese nombre.")
        return None, text
    item_id, canonical_name = matches[0]
    return item_id, canonical_name


def _stock_value(value, errors: list[str]) -> int | None:
    if _is_empty(value):
        errors.append("Stock disponible es obligatorio.")
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        errors.append("Stock disponible debe ser un número entero.")
        return None
    try:
        numeric = Decimal(str(value))
    except InvalidOperation:
        errors.append("Stock disponible debe ser un número entero.")
        return None
    if not numeric.is_finite() or numeric != numeric.to_integral_value():
        errors.append("Stock disponible no acepta decimales.")
        return None
    if numeric < 0:
        errors.append("Stock disponible no puede ser negativo.")
        return None
    return int(numeric)


def _price_value(value, errors: list[str]) -> float | None:
    if _is_empty(value):
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        errors.append("Precio debe ser un número sin símbolos ni separadores de miles.")
        return None
    try:
        numeric = Decimal(str(value))
    except InvalidOperation:
        errors.append("Precio debe ser un número válido.")
        return None
    if not numeric.is_finite():
        errors.append("Precio debe ser un número válido.")
        return None
    if numeric < 0:
        errors.append("Precio no puede ser negativo.")
        return None
    if numeric.as_tuple().exponent < -2:
        errors.append("Precio admite como máximo dos decimales.")
        return None
    if numeric > Decimal("9999999999.99"):
        errors.append("Precio supera el límite permitido.")
        return None
    return float(numeric)


def _choice_value(
    value,
    *,
    label: str,
    choices: dict[str, object],
    default_key: str,
    errors: list[str],
):
    if _is_empty(value):
        return choices[default_key]
    if not isinstance(value, str):
        errors.append(f"{label} debe ser texto.")
        return choices[default_key]
    key = value.strip().casefold()
    if key not in choices:
        allowed = ", ".join(choice.capitalize() for choice in choices)
        errors.append(f"{label} debe ser uno de estos valores: {allowed}.")
        return choices[default_key]
    return choices[key]


def _highlights_value(value, errors: list[str]) -> list[str]:
    if _is_empty(value):
        return []
    if not isinstance(value, str):
        errors.append("Destacados debe ser texto separado por punto y coma.")
        return []
    return [item.strip() for item in value.split(";") if item.strip()]


def _specs_value(value, errors: list[str]) -> dict[str, str]:
    if _is_empty(value):
        return {}
    if not isinstance(value, str):
        errors.append("Especificaciones debe contener pares clave=valor.")
        return {}

    specs: dict[str, str] = {}
    for position, raw_pair in enumerate(value.split(";"), start=1):
        pair = raw_pair.strip()
        if not pair:
            continue
        if "=" not in pair:
            errors.append(
                f"Especificaciones: el elemento {position} no contiene el signo '='."
            )
            continue
        key, spec_value = (part.strip() for part in pair.split("=", 1))
        if not key:
            errors.append(f"Especificaciones: el elemento {position} tiene la clave vacía.")
            continue
        if key in specs:
            errors.append(f"Especificaciones: la clave '{key}' está repetida.")
            continue
        specs[key] = spec_value
    return specs


def _row_preview(
    row_number: int,
    values: list,
    formula_columns: list[str],
    *,
    brand_index,
    category_index,
) -> dict:
    raw = {
        header: _json_value(value)
        for header, value in zip(PRODUCT_HEADERS, values, strict=True)
    }
    errors: list[str] = []
    warnings: list[str] = []
    if formula_columns:
        errors.append(
            "No se permiten fórmulas en: " + ", ".join(formula_columns) + "."
        )

    name = _text_value(
        values[0],
        label="Nombre",
        errors=errors,
        required=True,
        min_length=3,
        max_length=200,
    )
    sku = _text_value(
        values[1],
        label="SKU",
        errors=errors,
        required=True,
        min_length=2,
        max_length=60,
    )
    sku = normalize_sku(sku)
    part_number = _text_value(
        values[2], label="Número de parte", errors=errors, max_length=80
    )
    brand_id, brand_name = _resolve_catalog_name(
        values[3], label="Marca", index=brand_index, errors=errors
    )
    category_id, category_name = _resolve_catalog_name(
        values[4], label="Categoría", index=category_index, errors=errors
    )
    available_stock = _stock_value(values[5], errors)
    is_used = _choice_value(
        values[6],
        label="Condición",
        choices=CONDITION_VALUES,
        default_key="nuevo",
        errors=errors,
    )
    price = _price_value(values[7], errors)
    currency = _choice_value(
        values[8],
        label="Moneda",
        choices=CURRENCY_VALUES,
        default_key="pen",
        errors=errors,
    )
    status = _choice_value(
        values[9],
        label="Estado",
        choices=STATUS_VALUES,
        default_key="publicado",
        errors=errors,
    )
    short_description = _text_value(
        values[10], label="Descripción corta", errors=errors
    )
    description = _text_value(values[11], label="Descripción", errors=errors)
    highlights = _highlights_value(values[12], errors)
    specs = _specs_value(values[13], errors)

    normalized = {
        "name": name,
        "sku": sku,
        "brand_id": brand_id,
        "category_id": category_id,
        "part_number": part_number,
        "short_description": short_description,
        "description": description,
        "specs": specs,
        "highlights": highlights,
        "stock_note": STOCK_NOTE_DEFAULT,
        "price": price,
        "currency": currency,
        "available_stock": available_stock,
        "stock_type": STOCK_TYPE_DEFAULT,
        "is_used": is_used,
        "status": status,
    }
    return {
        "row_number": row_number,
        "sku": sku or str(raw.get("sku") or "").strip(),
        "name": name or str(raw.get("nombre") or "").strip(),
        "brand": brand_name,
        "category": category_name,
        "valid": False,
        "errors": errors,
        "warnings": warnings,
        "raw": raw,
        "normalized": normalized,
    }


def validate_product_workbook(
    contents: bytes,
    *,
    filename: str,
    brands: Iterable[tuple[int, str]],
    categories: Iterable[tuple[int, str]],
    existing_skus: Iterable[str],
) -> dict:
    if not contents:
        raise ProductImportFileError("El archivo está vacío.")
    if not contents.startswith(b"PK"):
        raise ProductImportFileError("El archivo no es un .xlsx válido o está corrupto.")

    try:
        workbook = load_workbook(
            BytesIO(contents), read_only=True, data_only=False, keep_links=False
        )
    except Exception as exc:
        raise ProductImportFileError(
            "El archivo no es un .xlsx válido o está corrupto."
        ) from exc

    try:
        if "Productos" not in workbook.sheetnames:
            raise ProductImportFileError("El archivo no contiene la hoja 'Productos'.")
        worksheet = workbook["Productos"]
        if worksheet.max_column != len(PRODUCT_HEADERS):
            raise ProductImportFileError(
                "La hoja Productos debe contener únicamente los 14 encabezados esperados."
            )

        actual_headers = [
            worksheet.cell(row=1, column=column).value
            for column in range(1, len(PRODUCT_HEADERS) + 1)
        ]
        if actual_headers != PRODUCT_HEADERS:
            raise ProductImportFileError(
                "Los encabezados de la hoja Productos son incorrectos o están fuera de orden."
            )

        brand_index = _build_name_index(brands)
        category_index = _build_name_index(categories)
        rows: list[dict] = []

        for cells in worksheet.iter_rows(
            min_row=2,
            max_row=worksheet.max_row,
            min_col=1,
            max_col=len(PRODUCT_HEADERS),
        ):
            values = [cell.value for cell in cells]
            if all(_is_empty(value) for value in values):
                continue
            row_number = cells[0].row
            if row_number > MAX_PRODUCT_ROW:
                raise ProductImportFileError(
                    "El archivo contiene datos después de la fila 300; se permiten como máximo 299 productos."
                )
            formula_columns = [
                PRODUCT_HEADERS[index]
                for index, cell in enumerate(cells)
                if cell.data_type == "f"
            ]
            rows.append(
                _row_preview(
                    row_number,
                    values,
                    formula_columns,
                    brand_index=brand_index,
                    category_index=category_index,
                )
            )

        if not rows:
            raise ProductImportFileError("El archivo no contiene filas de productos.")

        sku_counts = Counter(
            row["normalized"]["sku"]
            for row in rows
            if row["normalized"]["sku"]
        )
        catalog_skus = {normalize_sku(sku) for sku in existing_skus}
        for row in rows:
            sku = row["normalized"]["sku"]
            if sku and sku_counts[sku] > 1:
                row["errors"].append("SKU repetido dentro del archivo.")
            if sku and sku in catalog_skus:
                row["errors"].append("SKU ya existe en el catálogo.")
            row["valid"] = not row["errors"]

        valid_rows = sum(row["valid"] for row in rows)
        invalid_rows = len(rows) - valid_rows
        return {
            "filename": safe_upload_filename(filename),
            "total_rows": len(rows),
            "valid_rows": valid_rows,
            "invalid_rows": invalid_rows,
            "can_import": bool(rows) and invalid_rows == 0,
            "rows": rows,
        }
    finally:
        workbook.close()


def validate_product_workbook_for_db(
    db: Session,
    contents: bytes,
    *,
    filename: str,
) -> dict:
    """Validate against the same active catalogs and current SKUs used by imports."""
    brands = (
        db.query(models.Brand.id, models.Brand.name)
        .filter(models.Brand.is_active.is_(True))
        .all()
    )
    categories = (
        db.query(models.Category.id, models.Category.name)
        .filter(models.Category.is_active.is_(True))
        .all()
    )
    existing_skus = [sku for (sku,) in db.query(models.Product.sku).all()]
    return validate_product_workbook(
        contents,
        filename=filename,
        brands=brands,
        categories=categories,
        existing_skus=existing_skus,
    )


def create_products_from_preview(
    db: Session,
    user: models.AdminUser,
    preview: dict,
) -> list[dict]:
    """Persist a fully valid preview atomically and return compact product details."""
    if not preview.get("can_import") or not preview.get("rows"):
        raise ValueError("La vista previa contiene filas inválidas o está vacía.")

    imported: list[dict] = []
    try:
        for row in preview["rows"]:
            if not row.get("valid"):
                raise ValueError("La vista previa contiene filas inválidas.")
            payload = schemas.AdminProductInput.model_validate(row["normalized"])
            product = create_product_record(db, payload)
            audit_product_creation(db, user, product, source="excel_import")
            imported.append(
                {
                    "id": product.id,
                    "sku": product.sku,
                    "name": product.name,
                    "slug": product.slug,
                }
            )
        db.commit()
    except Exception:
        db.rollback()
        raise
    return imported
