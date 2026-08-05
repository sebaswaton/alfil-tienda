from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


PRODUCT_TEMPLATE_FILENAME = "plantilla_importacion_productos.xlsx"
PRODUCT_TEMPLATE_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
PRODUCT_HEADERS = [
    "nombre",
    "sku",
    "numero_parte",
    "marca",
    "categoria",
    "stock_disponible",
    "condicion",
    "precio",
    "moneda",
    "estado",
    "descripcion_corta",
    "descripcion",
    "destacados",
    "especificaciones",
]
PRODUCT_CONDITION_VALUES = ["Nuevo", "Usado"]
PRODUCT_CURRENCY_VALUES = ["PEN", "USD"]
PRODUCT_STATUS_VALUES = ["Publicado", "Borrador", "Oculto"]

HEADER_FILL = PatternFill("solid", fgColor="087F6F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SECTION_FILL = PatternFill("solid", fgColor="DFF3EF")
GRID_SIDE = Side(style="thin", color="AABAB7")
GRID_BORDER = Border(
    left=GRID_SIDE,
    right=GRID_SIDE,
    top=GRID_SIDE,
    bottom=GRID_SIDE,
)


def _inline_validation_formula(values: list[str]) -> str | None:
    """Return an inline Excel list when it fits the format's 255-char limit."""
    cleaned = [value.strip() for value in values if value and value.strip()]
    if not cleaned or any(
        "," in value or '"' in value or "\n" in value for value in cleaned
    ):
        return None
    formula = f'"{",".join(cleaned)}"'
    return formula if len(formula) <= 255 else None


def _add_list_validation(
    worksheet,
    cell_range: str,
    values: list[str],
    *,
    prompt: str,
) -> None:
    formula = _inline_validation_formula(values)
    if not formula:
        return
    validation = DataValidation(
        type="list",
        formula1=formula,
        allow_blank=True,
        error="Selecciona un valor de la lista.",
        errorTitle="Valor no permitido",
        prompt=prompt,
        promptTitle="Valores disponibles",
        showErrorMessage=True,
        showInputMessage=True,
    )
    worksheet.add_data_validation(validation)
    validation.add(cell_range)


def _build_products_sheet(workbook: Workbook, brands: list[str], categories: list[str]):
    worksheet = workbook.active
    worksheet.title = "Productos"
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = "A1:N301"
    worksheet.row_dimensions[1].height = 28

    for column, header in enumerate(PRODUCT_HEADERS, start=1):
        cell = worksheet.cell(row=1, column=column, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = GRID_BORDER

    widths = {
        "A": 36,
        "B": 24,
        "C": 22,
        "D": 24,
        "E": 27,
        "F": 21,
        "G": 17,
        "H": 18,
        "I": 15,
        "J": 18,
        "K": 36,
        "L": 50,
        "M": 44,
        "N": 56,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    for row in range(2, 302):
        worksheet.row_dimensions[row].height = 22
        for column in range(1, 15):
            worksheet.cell(row=row, column=column).border = GRID_BORDER
        worksheet[f"B{row}"].number_format = "@"
        worksheet[f"F{row}"].number_format = "0"
        worksheet[f"H{row}"].number_format = "0.00"
        for column in ("K", "L", "M", "N"):
            worksheet[f"{column}{row}"].alignment = Alignment(
                vertical="top", wrap_text=True
            )

    _add_list_validation(
        worksheet,
        "G2:G301",
        PRODUCT_CONDITION_VALUES,
        prompt="Selecciona Nuevo o Usado.",
    )
    _add_list_validation(
        worksheet,
        "I2:I301",
        PRODUCT_CURRENCY_VALUES,
        prompt="Selecciona PEN o USD.",
    )
    _add_list_validation(
        worksheet,
        "J2:J301",
        PRODUCT_STATUS_VALUES,
        prompt="Selecciona el estado del producto.",
    )
    _add_list_validation(
        worksheet,
        "D2:D301",
        brands,
        prompt="Selecciona una marca existente.",
    )
    _add_list_validation(
        worksheet,
        "E2:E301",
        categories,
        prompt="Selecciona una categoría existente.",
    )


def _build_instructions_sheet(workbook: Workbook):
    worksheet = workbook.create_sheet("Instrucciones")
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A5"

    worksheet.merge_cells("A1:D1")
    title = worksheet["A1"]
    title.value = "Plantilla de importación de productos"
    title.fill = HEADER_FILL
    title.font = Font(color="FFFFFF", bold=True, size=14)
    title.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 30

    worksheet.merge_cells("A2:D2")
    worksheet["A2"] = (
        "Completa una fila por producto. Esta plantilla replica los campos del "
        "formulario manual de productos."
    )
    worksheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    worksheet.row_dimensions[2].height = 34

    instruction_headers = ["Campo", "Obligatorio", "Formato o valores", "Indicaciones"]
    for column, value in enumerate(instruction_headers, start=1):
        cell = worksheet.cell(row=4, column=column, value=value)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    rows = [
        ("nombre", "Sí", "Texto", "Nombre comercial del producto."),
        ("sku", "Sí", "Texto", "Código interno que identifica el producto."),
        ("numero_parte", "No", "Texto", "Número de parte del fabricante."),
        ("marca", "Sí", "Lista o texto", "Nombre de una marca existente."),
        ("categoria", "Sí", "Lista o texto", "Nombre de una categoría existente."),
        ("stock_disponible", "Sí", "Entero desde 0", "Cantidad disponible."),
        ("condicion", "No", "Nuevo / Usado", "Si queda vacío se usará Nuevo."),
        (
            "precio",
            "No",
            "Decimal",
            'Déjalo vacío para mostrar "Cotizar". No escribas símbolos de moneda.',
        ),
        ("moneda", "No", "PEN / USD", "Si queda vacío se usará PEN."),
        (
            "estado",
            "No",
            "Publicado / Borrador / Oculto",
            "Si queda vacío se usará Publicado.",
        ),
        ("descripcion_corta", "No", "Texto largo", "Resumen para la tarjeta del catálogo."),
        ("descripcion", "No", "Texto largo", "Descripción completa del producto."),
        (
            "destacados",
            "No",
            "Lista separada por ;",
            "Ejemplo: Diseño compacto; Garantía del fabricante",
        ),
        (
            "especificaciones",
            "No",
            "Pares clave=valor separados por ;",
            "Ejemplo: Procesador=Intel Core i5; Memoria=16 GB",
        ),
    ]
    for row_number, values in enumerate(rows, start=5):
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row_number, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    notes_row = 20
    worksheet.merge_cells(
        start_row=notes_row, start_column=1, end_row=notes_row, end_column=4
    )
    notes_title = worksheet.cell(
        row=notes_row, column=1, value="Indicaciones importantes"
    )
    notes_title.fill = SECTION_FILL
    notes_title.font = Font(bold=True, color="075E54")

    notes = [
        "No modifiques, elimines ni reordenes los encabezados de la hoja Productos.",
        "Cada fila de la hoja Productos representa un producto.",
        "Las filas completamente vacías serán ignoradas en la futura importación.",
        "En destacados, separa cada elemento con punto y coma (;).",
        "En especificaciones, usa clave=valor y separa los pares con punto y coma (;).",
    ]
    for row_number, note in enumerate(notes, start=notes_row + 1):
        worksheet.merge_cells(
            start_row=row_number, start_column=1, end_row=row_number, end_column=4
        )
        cell = worksheet.cell(row=row_number, column=1, value=f"• {note}")
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for column, width in {"A": 24, "B": 16, "C": 31, "D": 72}.items():
        worksheet.column_dimensions[column].width = width


def build_product_import_template(
    *, brands: list[str] | None = None, categories: list[str] | None = None
) -> BytesIO:
    workbook = Workbook()
    _build_products_sheet(workbook, brands or [], categories or [])
    _build_instructions_sheet(workbook)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
