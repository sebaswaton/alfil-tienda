import unittest

from openpyxl import load_workbook

from app.product_excel_template import PRODUCT_HEADERS, build_product_import_template


class ProductExcelTemplateTests(unittest.TestCase):
    def setUp(self):
        contents = build_product_import_template(
            brands=["HP", "Cisco"],
            categories=["Laptops", "Redes y Conectividad"],
        )
        self.workbook = load_workbook(contents)

    def test_workbook_structure_and_headers(self):
        self.assertEqual(self.workbook.sheetnames, ["Productos", "Instrucciones"])
        self.assertTrue(
            all(sheet.sheet_state == "visible" for sheet in self.workbook)
        )

        products = self.workbook["Productos"]
        headers = [
            products.cell(row=1, column=column).value for column in range(1, 15)
        ]
        self.assertEqual(headers, PRODUCT_HEADERS)
        self.assertEqual(products.max_column, 14)
        self.assertEqual(products.max_row, 301)
        self.assertEqual(products.freeze_panes, "A2")
        self.assertEqual(products.auto_filter.ref, "A1:N301")
        self.assertTrue(
            all(
                products.cell(row=row, column=column).value is None
                for row in range(2, 302)
                for column in range(1, 15)
            )
        )
        self.assertTrue(
            all(
                products.cell(row=row, column=column).border.left.style == "thin"
                and products.cell(row=row, column=column).border.right.style == "thin"
                and products.cell(row=row, column=column).border.top.style == "thin"
                and products.cell(row=row, column=column).border.bottom.style == "thin"
                for row in range(1, 302)
                for column in range(1, 15)
            )
        )
        self.assertFalse(
            any(
                dimension.hidden
                for dimension in products.column_dimensions.values()
            )
        )
        self.assertEqual(
            self.workbook["Instrucciones"]["A1"].value,
            "Plantilla de importación de productos",
        )

    def test_number_formats_and_dropdowns(self):
        products = self.workbook["Productos"]
        self.assertEqual(products["F2"].number_format, "0")
        self.assertEqual(products["H2"].number_format, "0.00")

        validation_ranges = {
            str(validation.sqref): validation.formula1
            for validation in products.data_validations.dataValidation
        }
        self.assertEqual(validation_ranges["G2:G301"], '"Nuevo,Usado"')
        self.assertEqual(validation_ranges["I2:I301"], '"PEN,USD"')
        self.assertEqual(
            validation_ranges["J2:J301"], '"Publicado,Borrador,Oculto"'
        )
        self.assertEqual(validation_ranges["D2:D301"], '"HP,Cisco"')
        self.assertEqual(
            validation_ranges["E2:E301"], '"Laptops,Redes y Conectividad"'
        )


if __name__ == "__main__":
    unittest.main()
